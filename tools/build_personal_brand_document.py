"""Create an editable website content brief using the bundled document runtime."""
import json
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT.parent / 'Vivek_Pandey_Website'
OUT.mkdir(exist_ok=True)
PAGES = json.loads((ROOT / 'website/personal-brand-pages.json').read_text())
BASE = 'https://academy-tools-jade.vercel.app'
doc = Document()
section = doc.sections[0]
section.page_width=Inches(8.5); section.page_height=Inches(11)
section.top_margin=Inches(.7); section.bottom_margin=Inches(.7)
section.left_margin=Inches(.8); section.right_margin=Inches(.8)
section.footer_distance=Inches(.3)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    style=doc.styles[name]; style.font.name='Arial'; style.font.color.rgb=RGBColor(0,0,0)
    style.paragraph_format.space_after=Pt(5)
    if style.element.pPr is not None:
        for border in list(style.element.pPr.findall(qn('w:pBdr'))):
            style.element.pPr.remove(border)
doc.styles['Normal'].font.size=Pt(11)
doc.styles['Normal'].paragraph_format.line_spacing=1.08
doc.styles['Title'].font.size=Pt(25)
doc.styles['Title'].paragraph_format.space_after=Pt(14)
doc.styles['Heading 1'].font.size=Pt(23)
doc.styles['Heading 1'].paragraph_format.space_before=Pt(0)
doc.styles['Heading 1'].paragraph_format.space_after=Pt(12)
doc.styles['Heading 2'].font.size=Pt(12)
doc.styles['Heading 2'].paragraph_format.space_before=Pt(11)
doc.styles['Heading 2'].paragraph_format.space_after=Pt(5)
doc.styles['Heading 3'].font.size=Pt(11)
doc.styles['Heading 3'].paragraph_format.space_before=Pt(5)
doc.styles['Heading 3'].paragraph_format.space_after=Pt(3)
doc.core_properties.title='Vivek Pandey website content and direction'
doc.core_properties.subject='Personal website and six public pages'
doc.core_properties.author='Vivek Pandey'
doc.core_properties.keywords='AI consulting, AI speaking, AI courses, website'

def p(text): return doc.add_paragraph(text)
def h(text): return doc.add_heading(text,2)
def hyperlink(paragraph,label,url):
    rid=paragraph.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    node=OxmlElement('w:hyperlink'); node.set(qn('r:id'),rid)
    run=OxmlElement('w:r'); props=OxmlElement('w:rPr')
    col=OxmlElement('w:color'); col.set(qn('w:val'),'305744'); props.append(col)
    underline=OxmlElement('w:u'); underline.set(qn('w:val'),'single'); props.append(underline)
    run.append(props); text=OxmlElement('w:t'); text.text=label; run.append(text)
    node.append(run); paragraph._p.append(node)
def page(title,slug):
    title_row=doc.add_heading(title,1)
    title_row.paragraph_format.page_break_before=True
    row=p('Live page  '); hyperlink(row,BASE+slug,BASE+slug)
    row.paragraph_format.space_after=Pt(14)
def labelled(label,text):
    row=p(''); row.add_run(label+'  ').bold=True; row.add_run(text)

doc.add_paragraph('Vivek Pandey website content and direction','Title')
p('The website presents Vivek Pandey as an independent personal brand for AI consulting, speaking and education. It should help a business leader, event organiser or learner understand the offer and take a clear next step.')
h('Purpose and audience')
p('The aim is to build a credible professional presence for international clients while making the services easy to understand. The main audience includes business owners and leaders, software teams, manufacturing organisations, universities, colleges and conference organisers. The profile highlights 19 years of industry experience.')
h('The three business areas')
labelled('AI consulting','Help businesses understand where AI fits, compare solutions and plan implementation. Manufacturing, software development, sports and apparel provide starting points for discussion.')
labelled('AI speaking','Explain practical AI applications through talks and workshops shaped around the audience and the purpose of the event.')
labelled('AI courses','Offer advanced learning through the personal website, with Wooplix Academy providing the existing course library. Training enquiries are handled separately from the reading material.')
h('Website structure and live pages')
table=doc.add_table(rows=1,cols=2); table.alignment=WD_TABLE_ALIGNMENT.CENTER
table.autofit=False; table.columns[0].width=Inches(1.45); table.columns[1].width=Inches(5.45)
for cell,text in zip(table.rows[0].cells,['Page','Live address']):
    cell.text=text; cell.paragraphs[0].runs[0].bold=True;cell.paragraphs[0].runs[0].font.color.rgb=RGBColor(255,255,255)
    shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'414141');cell._tc.get_or_add_tcPr().append(shading)
repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(repeat)
routes=[('Home','/'),('About','/about'),('Contact','/contact'),('AI Consulting','/ai-consulting'),('AI Speaking','/ai-speaking'),('AI Courses','/ai-courses')]
for label,route in routes:
    a,b=table.add_row().cells;a.text=label;hyperlink(b.paragraphs[0],BASE+route,BASE+route)
for row in table.rows:
    for cell in row.cells:
        cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        tcpr=cell._tc.get_or_add_tcPr(); borders=OxmlElement('w:tcBorders')
        for edge in ['top','left','bottom','right']:
            e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
        tcpr.append(borders)
        margins=OxmlElement('w:tcMar')
        for edge in ['top','left','bottom','right']:
            e=OxmlElement('w:'+edge);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');margins.append(e)
        tcpr.append(margins)
        for para in cell.paragraphs:
            para.paragraph_format.space_after=Pt(2);para.paragraph_format.space_before=Pt(2)
h('The public address and booking approach')
p('The intended domain is vivekpandey.ai. The pages currently run at the Vercel address above. Connecting the custom domain is a separate launch step. Visitors can request 30-minute or 60-minute consulting sessions. Availability, scope and fees are agreed before confirmation.')

page('Home','/')
h('Main message')
p('Make AI work for your business.')
p('Find the right opportunity. Build a practical plan. Give your people the skills to carry it forward.')
p('For business leaders, technical teams and organisations exploring their next move with AI.')
h('The opening question')
p('What should AI change in your business? A useful answer connects the technology to your process, people, data and responsibilities.')
h('Services visitors can explore')
labelled('Consulting','Understand where AI fits, what it depends on and how to turn an idea into a focused pilot.')
labelled('Speaking','Practical talks and workshops that connect AI concepts to decisions people make at work.')
labelled('Courses','Structured learning in AI, agents, business automation and implementation through the Academy.')
h('Experience and approach')
p('With 19 years of industry experience, I focus on a straightforward question: how can technology make business work better? My work brings together AI consulting, speaking and education.')
p('For consulting, the process begins with the current work, compares possible solutions and defines a pilot before planning delivery and adoption. Example discussion areas include manufacturing, software teams, sports and apparel.')
h('Learning and future development')
p('The Academy remains the education part of the website. Visitors can open the AI Builder programme or Corporate AI and Automation Workshop, then browse all five programmes and their Course Books.')
p('Research directions include an Agentic AI Lab and software delivery supported by specialist agents. Delivery speed and business results need measurement in real projects.')
h('The next step')
p('Visitors can choose a 30-minute focused discussion or a 60-minute working discussion. The request page asks for the business question and relevant context. The Home page also links directly to About, Contact and each service page.')

for slug in ['ai-consulting','ai-speaking','ai-courses','about']:
    data=PAGES[slug]; page(data['title'],'/'+slug)
    p(data['headline']);p(data['intro'])
    for item in data['sections']:
        h(item['title'])
        if item.get('intro'):p(item['intro'])
        for card in item.get('cards',[]):
            labelled(card['title'],card['text'])
        items=item.get('items',[])
        if slug=='ai-consulting' and item['title']=='How a consulting engagement works':
            items=[
                'Walk through the process, the people involved and the problem that needs attention.',
                'Compare existing tools, automation and AI, including their dependencies and limitations.',
                'Agree a pilot scope, access, measures, review responsibilities and conditions for continuing.',
                'Plan implementation, training and support. Agree separate deliverables for a larger engagement.'
            ]
        for i,text in enumerate(items,1):p(f'{i}. {text}')
        for text in item.get('paragraphs',[]):p(text)
    if slug=='ai-consulting':
        h('Consulting session options')
        labelled('30 minutes','A focused discussion of one opportunity or decision, its constraints and a useful next step.')
        labelled('60 minutes','A deeper discussion of a process, possible approaches, risks and priorities for a pilot.')
        p('Both are session requests. Timing, scope and fees are agreed before confirmation.')
    if slug=='ai-courses':
        row=p('Complete course library  ');hyperlink(row,BASE+'/academy',BASE+'/academy')
    labelled('Main action',data['cta'])

page('Contact','/contact')
h('Main message')
p('Contact Vivek Pandey.')
p('A little context makes the conversation more useful. Prepare your enquiry here, then send it from your email app.')
h('How to get in touch')
p('Use the contact page for consulting, speaking, course or general enquiries. Email directly at wooplix15@gmail.com if that is more convenient. Include your location or time zone so arrangements can be discussed clearly.')
h('The enquiry form')
p('The form asks for your name, email, organisation, enquiry type, preferred session, time zone and a short question or brief. Name, email and the brief are required. The session choices are 30 minutes, 60 minutes or not specified.')
p('The Open email draft button prepares an email to wooplix15@gmail.com in your email app. Review it before sending. Copy enquiry provides another way to paste the details into an email. The website does not send the enquiry automatically or reserve a calendar slot.')
h('What to include')
labelled('Consulting','Your industry, the current process and systems, the main challenge and the decision you need help making.')
labelled('Speaking','The audience, event purpose, date, location or online format, preferred topic and session format.')
labelled('Courses','Your current experience, the programme that interests you and the skill you or your team need to develop.')
h('Session requests')
p('A 30-minute option suits a focused question. A 60-minute option allows more time to examine a process or implementation idea. Timing, scope and fees are confirmed separately. Use a brief description rather than confidential client documents for the first enquiry.')
h('Visitor actions')
p('Open email draft. Copy enquiry. Email directly. Return to the Home page or explore the Academy.')

footer=section.footer.paragraphs[0]
footer.paragraph_format.space_after=Pt(0)
run=footer.add_run('Vivek Pandey   |   Website content   |   ');run.font.size=Pt(9)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
output=OUT/'Vivek_Pandey_Website_Content_and_Plan.docx'
doc.save(output)
print(output)
